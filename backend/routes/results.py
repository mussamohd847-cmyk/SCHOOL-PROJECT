from flask import Blueprint, request
from openpyxl import load_workbook

from database import db
from models import Result, Exam, Student, Subject
from utils.auth import login_required, roles_required
from utils.responses import (
    success_response,
    error_response,
    paginate_query,
)
from utils.validators import require_fields


results_bp = Blueprint(
    "results",
    __name__,
    url_prefix="/api/results",
)


def calculate_total(hw, ct, cw, fe):
    hw = float(hw or 0)
    ct = float(ct or 0)
    cw = float(cw or 0)
    fe = float(fe or 0)

    total = (
        hw * 0.10
        + ct * 0.20
        + cw * 0.10
        + fe * 0.60
    )

    return round(total, 2)


def calculate_grade(total):
    total = float(total)

    if total >= 80:
        return "A"

    if total >= 65:
        return "B"

    if total >= 50:
        return "C"

    if total >= 35:
        return "D"

    return "F"


@results_bp.route("", methods=["GET"])
@login_required
def list_results():
    query = Result.query

    for field in (
        "student_id",
        "exam_id",
        "subject_id",
    ):
        value = request.args.get(
            field,
            type=int,
        )

        if value:
            query = query.filter(
                getattr(Result, field) == value
            )

    items, pagination = paginate_query(
        query,
        request.args,
    )

    return success_response(
        "Results list",
        [r.to_dict() for r in items],
        pagination=pagination,
    )


@results_bp.route("", methods=["POST"])
@login_required
@roles_required("ADMIN", "TEACHER")
def create_result():
    data = request.get_json(
        silent=True
    ) or {}

    missing = require_fields(
        data,
        [
            "student_id",
            "exam_id",
            "subject_id",
            "fe",
        ],
    )

    if missing:
        return error_response(
            f"{missing} is required",
            400,
        )

    student_id = data["student_id"]
    exam_id = data["exam_id"]
    subject_id = data["subject_id"]

    existing = Result.query.filter_by(
        student_id=student_id,
        exam_id=exam_id,
        subject_id=subject_id,
    ).first()

    if existing:
        return error_response(
            "A result for this student/exam/subject already exists",
            409,
        )

    student = Student.query.get(
        student_id
    )

    if not student:
        return error_response(
            "Student not found",
            404,
        )

    exam = Exam.query.get(
        exam_id
    )

    if not exam:
        return error_response(
            "Exam not found",
            404,
        )

    try:
        hw = float(
            data.get("hw", 0)
        )

        ct = float(
            data.get("ct", 0)
        )

        cw = float(
            data.get("cw", 0)
        )

        fe = float(
            data.get("fe", 0)
        )
    except (
        TypeError,
        ValueError,
    ):
        return error_response(
            "Marks must be valid numbers",
            400,
        )

    for value in (
        hw,
        ct,
        cw,
        fe,
    ):
        if value < 0 or value > 100:
            return error_response(
                "Each mark must be between 0 and 100",
                400,
            )

    total = calculate_total(
        hw,
        ct,
        cw,
        fe,
    )

    grade = calculate_grade(
        total
    )

    result = Result(
        student_id=student_id,
        exam_id=exam_id,
        subject_id=subject_id,
        class_id=data.get("class_id"),
        stream_id=data.get("stream_id"),
        hw=hw,
        ct=ct,
        cw=cw,
        fe=fe,
        total=total,
        grade=grade,
    )

    db.session.add(result)
    db.session.commit()

    return success_response(
        "Result recorded successfully",
        result.to_dict(),
        201,
    )


@results_bp.route(
    "/import-excel",
    methods=["POST"],
)
@login_required
@roles_required("ADMIN", "TEACHER")
def import_results_excel():
    exam_id = request.form.get("exam_id", type=int)
    subject_id = request.form.get("subject_id", type=int)
    file = request.files.get("file")

    if not exam_id:
        return error_response("exam_id is required", 400)

    if not subject_id:
        return error_response("subject_id is required", 400)

    if not file:
        return error_response("Excel file is required", 400)

    if not file.filename:
        return error_response("Please select an Excel file", 400)

    if not file.filename.lower().endswith((".xlsx", ".xlsm")):
        return error_response(
            "Only .xlsx or .xlsm Excel files are supported",
            400,
        )

    exam = Exam.query.get(exam_id)

    if not exam:
        return error_response("Exam not found", 404)

    subject = Subject.query.get(subject_id)

    if not subject:
        return error_response("Subject not found", 404)

    if subject.category != "MADRASA":
        return error_response(
            "Only MADRASA subjects can be imported from this page",
            400,
        )

    try:
        workbook = load_workbook(
            file,
            read_only=True,
            data_only=True,
        )
        worksheet = workbook.active
        rows = worksheet.iter_rows(values_only=True)
        headers = next(rows)
    except Exception as exc:
        return error_response(
            f"Unable to read Excel file: {str(exc)}",
            400,
        )

    normalized_headers = [
        str(header).strip().lower().replace(" ", "_")
        if header is not None else ""
        for header in headers
    ]

    required_headers = {
        "admission_no",
        "hw",
        "ct",
        "cw",
        "fe",
    }

    missing_headers = required_headers.difference(
        normalized_headers
    )

    if missing_headers:
        return error_response(
            "Missing Excel columns: "
            + ", ".join(sorted(missing_headers)),
            400,
        )

    header_index = {
        header: index
        for index, header in enumerate(normalized_headers)
        if header
    }

    imported = []
    skipped = []
    errors = []

    for row_number, row in enumerate(rows, start=2):
        try:
            admission_value = row[
                header_index["admission_no"]
            ]

            if admission_value is None:
                continue

            admission_no = str(
                admission_value
            ).strip()

            if not admission_no:
                continue

            student = Student.query.filter_by(
                admission_no=admission_no
            ).first()

            if not student:
                errors.append({
                    "row": row_number,
                    "admission_no": admission_no,
                    "error": "Student not found",
                })
                continue

            if not student.in_madrasa:
                errors.append({
                    "row": row_number,
                    "admission_no": admission_no,
                    "error": "Student is not enrolled in Madrasa",
                })
                continue

            if not student.madrasa_class_id:
                errors.append({
                    "row": row_number,
                    "admission_no": admission_no,
                    "error": "Student has no Madrasa class",
                })
                continue

            def number(column):
                value = row[header_index[column]]

                if value is None or value == "":
                    return 0.0

                return float(value)

            hw = number("hw")
            ct = number("ct")
            cw = number("cw")
            fe = number("fe")

            for mark in (hw, ct, cw, fe):
                if mark < 0 or mark > 100:
                    raise ValueError(
                        "Each mark must be between 0 and 100"
                    )

            existing = Result.query.filter_by(
                student_id=student.id,
                exam_id=exam_id,
                subject_id=subject_id,
            ).first()

            if existing:
                skipped.append({
                    "row": row_number,
                    "admission_no": admission_no,
                    "reason": "Result already exists",
                })
                continue

            total = calculate_total(
                hw,
                ct,
                cw,
                fe,
            )

            grade = calculate_grade(total)

            result = Result(
                student_id=student.id,
                exam_id=exam_id,
                subject_id=subject_id,
                class_id=student.madrasa_class_id,
                stream_id=student.madrasa_stream_id,
                hw=hw,
                ct=ct,
                cw=cw,
                fe=fe,
                total=total,
                grade=grade,
            )

            db.session.add(result)

            imported.append({
                "row": row_number,
                "admission_no": admission_no,
                "total": total,
                "grade": grade,
            })

        except (TypeError, ValueError) as exc:
            errors.append({
                "row": row_number,
                "admission_no": (
                    str(
                        row[
                            header_index["admission_no"]
                        ]
                    ).strip()
                    if row[
                        header_index["admission_no"]
                    ] is not None
                    else ""
                ),
                "error": str(exc),
            })

    try:
        db.session.commit()
    except Exception as exc:
        db.session.rollback()
        return error_response(
            f"Failed to save imported results: {str(exc)}",
            500,
        )

    return success_response(
        "Excel results import completed",
        {
            "exam_id": exam_id,
            "subject_id": subject_id,
            "imported_count": len(imported),
            "skipped_count": len(skipped),
            "error_count": len(errors),
            "imported": imported,
            "skipped": skipped,
            "errors": errors,
        },
    )


@results_bp.route(
    "/<int:result_id>",
    methods=["PUT"],
)
@login_required
@roles_required("ADMIN", "TEACHER")
def update_result(result_id):
    result = Result.query.get(
        result_id
    )

    if not result:
        return error_response(
            "Result not found",
            404,
        )

    data = request.get_json(
        silent=True
    ) or {}

    try:
        if "hw" in data:
            result.hw = float(
                data["hw"]
            )

        if "ct" in data:
            result.ct = float(
                data["ct"]
            )

        if "cw" in data:
            result.cw = float(
                data["cw"]
            )

        if "fe" in data:
            result.fe = float(
                data["fe"]
            )
    except (
        TypeError,
        ValueError,
    ):
        return error_response(
            "Marks must be valid numbers",
            400,
        )

    for value in (
        result.hw or 0,
        result.ct or 0,
        result.cw or 0,
        result.fe or 0,
    ):
        if value < 0 or value > 100:
            return error_response(
                "Each mark must be between 0 and 100",
                400,
            )

    result.total = calculate_total(
        result.hw,
        result.ct,
        result.cw,
        result.fe,
    )

    result.grade = calculate_grade(
        result.total
    )

    db.session.commit()

    return success_response(
        "Result updated successfully",
        result.to_dict(),
    )


@results_bp.route(
    "/<int:result_id>",
    methods=["DELETE"],
)
@login_required
@roles_required("ADMIN")
def delete_result(result_id):
    result = Result.query.get(
        result_id
    )

    if not result:
        return error_response(
            "Result not found",
            404,
        )

    db.session.delete(result)
    db.session.commit()

    return success_response(
        "Result deleted successfully"
    )


@results_bp.route(
    "/student/<int:student_id>/report",
    methods=["GET"],
)
@login_required
def student_report(student_id):
    results = Result.query.filter_by(
        student_id=student_id
    ).all()

    average = (
        round(
            sum(
                float(r.total or 0)
                for r in results
            )
            / len(results),
            2,
        )
        if results
        else 0
    )

    return success_response(
        "Student report",
        {
            "student_id": student_id,
            "average_marks": average,
            "results": [
                r.to_dict()
                for r in results
            ],
        },
    )


@results_bp.route(
    "/class/<int:class_id>/ranking",
    methods=["GET"],
)
@login_required
def class_ranking(class_id):
    exam_id = request.args.get(
        "exam_id",
        type=int,
    )

    students = Student.query.filter_by(
        class_id=class_id
    ).all()

    student_ids = [
        s.id
        for s in students
    ]

    if not student_ids:
        return success_response(
            "Class ranking",
            [],
        )

    query = Result.query.filter(
        Result.student_id.in_(
            student_ids
        )
    )

    if exam_id:
        query = query.filter_by(
            exam_id=exam_id
        )

    results = query.all()

    totals = {}

    for result in results:
        totals.setdefault(
            result.student_id,
            [],
        ).append(
            float(result.total or 0)
        )

    ranking = sorted(
        (
            {
                "student_id": student_id,
                "average_marks": round(
                    sum(values)
                    / len(values),
                    2,
                ),
            }
            for student_id, values
            in totals.items()
        ),
        key=lambda item:
            item["average_marks"],
        reverse=True,
    )

    for index, entry in enumerate(
        ranking,
        start=1,
    ):
        entry["rank"] = index

    return success_response(
        "Class ranking",
        ranking,
    )


@results_bp.route(
    "/subject/<int:subject_id>/performance",
    methods=["GET"],
)
@login_required
def subject_performance(subject_id):
    results = Result.query.filter_by(
        subject_id=subject_id
    ).all()

    if not results:
        return success_response(
            "Subject performance",
            {
                "subject_id": subject_id,
                "average_marks": 0,
                "count": 0,
                "grade_distribution": {},
            },
        )

    average = round(
        sum(
            float(r.total or 0)
            for r in results
        )
        / len(results),
        2,
    )

    grade_distribution = {}

    for result in results:
        grade_distribution[
            result.grade
        ] = (
            grade_distribution.get(
                result.grade,
                0,
            )
            + 1
        )

    return success_response(
        "Subject performance",
        {
            "subject_id": subject_id,
            "average_marks": average,
            "count": len(results),
            "grade_distribution":
                grade_distribution,
        },
    )